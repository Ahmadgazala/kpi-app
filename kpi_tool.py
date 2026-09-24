#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
KPI CLI Tool - Offline Excel updater
Usage:
  python3 kpi_tool.py --help
  python3 kpi_tool.py --interactive
  python3 kpi_tool.py --scores scores.json --out KPI-result.xlsx
  python3 kpi_tool.py --random --out demo.xlsx
"""
import argparse, json, pathlib, random
import openpyxl

BASE = pathlib.Path(__file__).parent
EXCEL = BASE / "KPI-Final 2026.xlsx"
JSON_DOMAINS = BASE / "kpi_domains.json"

LEVELS = {1:"ناشئ",2:"أولي",3:"أساسي",4:"منظم",5:"ديناميكي"}
def get_level(avg):
    if avg < 1.8: return "ناشئ"
    if avg < 2.6: return "أولي"
    if avg < 3.4: return "أساسي"
    if avg < 4.2: return "منظم"
    return "ديناميكي"

def load_domains():
    return json.loads(JSON_DOMAINS.read_text(encoding="utf-8"))

def build_scores_interactive(domains):
    scores = {}
    print("=== KPI 2026 - إدخال تفاعلي (1-5) ===\nاكتب 1-5 لكل مؤشر، أو Enter للتخطي (يبقى 1)\n")
    for d in domains:
        print(f"\n--- {d['domain']} ({len(d['indicators'])} مؤشرات) ---")
        scores[d['domain']] = {}
        for ind in d['indicators']:
            print(f"\n{ind['id']}. {ind['name']}")
            print(f"   {ind['question']}")
            for i, lv in enumerate(ind['levels'],1):
                print(f"   {i}={LEVELS[i]}: {lv[:70]}...")
            while True:
                ans = input(f"   درجتك [1-5, enter=1]: ").strip()
                if ans=="": val=1; break
                if ans in "12345": val=int(ans); break
                print("   أدخل 1-5 فقط")
            scores[d['domain']][ind['name']] = val
    return scores

def calculate(scores, domains):
    results=[]
    all_vals=[]
    for d in domains:
        vals=list(scores[d['domain']].values())
        avg=sum(vals)/len(vals)
        results.append({"domain":d['domain'],"avg":round(avg,2),"score100":round(avg*20,1),"level":get_level(avg),"count":len(vals)})
        all_vals.extend(vals)
    overall=sum(all_vals)/len(all_vals)
    return results, round(overall,2), round(overall*20,1), get_level(overall)

def update_excel(scores, out_path):
    domains=load_domains()
    wb=openpyxl.load_workbook(EXCEL)
    results, overall_avg, overall_100, overall_level = calculate(scores, domains)
    res_map={r["domain"]:r for r in results}
    # update domain sheets
    for d in domains:
        ws=wb[d["sheet"]]
        # find header
        hdr=None
        for idx,row in enumerate(ws.iter_rows(values_only=True),1):
            if row and any("1 - ناشئ" in str(c) for c in row if c):
                hdr=idx; break
        if not hdr: continue
        r=hdr+1
        for ind in d["indicators"]:
            sc=scores[d["domain"]][ind["name"]]
            ws.cell(row=r,column=8,value=sc)
            ws.cell(row=r,column=9,value=sc*20)
            r+=1
        # متوسط المجال
        for row in ws.iter_rows(min_row=r,max_row=ws.max_row):
            for cell in row:
                if cell.value and "متوسط المجال" in str(cell.value):
                    avg=sum(scores[d["domain"]].values())/len(d["indicators"])
                    ws.cell(row=cell.row,column=cell.column+1,value=round(avg,2))
                    ws.cell(row=cell.row,column=cell.column+2,value=round(avg*20,1))
    # النتائج
    if "النتائج" in wb.sheetnames:
        ws=wb["النتائج"]
        for row in ws.iter_rows(min_row=5,max_row=15):
            dc=row[0].value
            if dc and str(dc).strip():
                # fuzzy match
                for k,v in res_map.items():
                    if str(dc).strip() in k or k in str(dc).strip() or str(dc).strip()==k.strip():
                        row[2].value=v["avg"]
                        row[3].value=v["level"]
                        row[4].value=v["score100"]
                        break
                # handle abbreviated names
                if "الترويج" in str(dc): 
                    m=res_map.get("الترويج والخبرة في تعلم الكبار")
                    if m: row[2].value=m["avg"]; row[3].value=m["level"]; row[4].value=m["score100"]
        for row in ws.iter_rows(min_row=14,max_row=17):
            if row[0].value and "المتوسط العام" in str(row[0].value):
                row[2].value=overall_avg
                row[3].value=overall_level
                row[4].value=overall_100
    wb.save(out_path)
    return results, overall_avg, overall_100, overall_level

def main():
    p=argparse.ArgumentParser(description="KPI 2026 CLI")
    p.add_argument("--interactive", action="store_true", help="إدخال تفاعلي في الطرفية")
    p.add_argument("--scores", type=str, help="ملف JSON للدرجات")
    p.add_argument("--out", type=str, default="KPI-2026-result.xlsx", help="ملف الإخراج")
    p.add_argument("--random", action="store_true", help="توليد درجات عشوائية للتجربة")
    p.add_argument("--show", action="store_true", help="عرض النتائج فقط")
    args=p.parse_args()
    domains=load_domains()
    scores=None
    if args.scores:
        scores=json.loads(pathlib.Path(args.scores).read_text(encoding="utf-8"))
        # handle wrapped format
        if "scores" in scores: scores=scores["scores"]
    elif args.random:
        scores={d["domain"]:{ind["name"]: random.randint(1,5) for ind in d["indicators"]} for d in domains}
        pathlib.Path("scores-demo.json").write_text(json.dumps(scores,ensure_ascii=False,indent=2),encoding="utf-8")
        print("تم توليد scores-demo.json")
    elif args.interactive:
        scores=build_scores_interactive(domains)
        pathlib.Path("scores.json").write_text(json.dumps(scores,ensure_ascii=False,indent=2),encoding="utf-8")
        print("تم حفظ scores.json")
    else:
        # default: all 1
        scores={d["domain"]:{ind["name"]:1 for ind in d["indicators"]} for d in domains}
        print("No input -> using all 1 (ناشئ). Use --interactive or --scores")
    results, avg, score100, level = calculate(scores, domains)
    print("\n=== النتائج ===")
    for r in results:
        print(f"{r['domain']:35} | {r['avg']:.2f} | {r['level']:6} | {r['score100']:.0f}/100")
    print(f"\nالمتوسط العام: {avg:.2f} | {level} | {score100:.0f}/100")
    if not args.show:
        out=pathlib.Path(args.out)
        update_excel(scores, out)
        print(f"\n✓ تم إنشاء: {out.resolve()}")

if __name__=="__main__":
    main()
